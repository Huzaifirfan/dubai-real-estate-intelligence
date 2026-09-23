"""Stage 9: build the Excel executive workbook from the protected Stage 6 CSV.

Run from any directory with the project's existing virtual environment:
    .venv/bin/python src/09_build_excel_workbook.py

Only this stage's workbook and report are written. No database is contacted.
The full CSV supplies every analysis; the 5,000-row sample is for inspection.
"""

import ast
import calendar
from collections import Counter
from datetime import datetime
from decimal import Decimal
import hashlib
import math
from pathlib import Path
import tempfile
import warnings
import zipfile
import xml.etree.ElementTree as ET

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.chart.data_source import AxDataSource, StrRef, StrData, StrVal, NumData, NumVal
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.comments import Comment
from openpyxl.drawing.text import Paragraph, ParagraphProperties, CharacterProperties, RegularTextRun
from openpyxl.chart.text import RichText
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.views import Selection
from openpyxl.workbook.properties import CalcProperties


ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / 'data/processed/dld_transactions_2026_ytd_features.csv'
OUTPUT = ROOT / 'excel/Dubai_Real_Estate_Intelligence_2026_YTD.xlsx'
REPORT = ROOT / 'reports/stage9_excel_workbook_report.txt'
SHEETS = ['Executive Summary', 'Monthly Analysis', 'Area Analysis',
          'Segment Analysis', 'Data Sample', 'Data Dictionary', 'Notes']
SAMPLE_SIZE = 5000
SAMPLE_SEED = 20260921
NAVY, TEAL, GOLD = '183449', '137C83', 'C89339'
INK, MUTED, LIGHT, LINE = '203547', '566A78', 'EFF5F7', 'DCE5EA'
COUNT_FMT = '#,##0'
AED_FMT = '"AED "#,##0.00'
VALUE_FMT = '"AED "#,##0'
PCT_FMT = '0.00%'
DATE_FMT = 'dd-mmm-yyyy'
DATETIME_FMT = 'dd-mmm-yyyy hh:mm:ss'
VALUE_BANDS = ['Under 500K', '500K - 1M', '1M - 2M', '2M - 5M', '5M - 10M', '10M+']
SIZE_BANDS = ['Under 50 sqm', '50 - 100 sqm', '100 - 200 sqm', '200 - 500 sqm', '500+ sqm']
PRICE_NOTE = ('Price statistics: Sales with VALID_SALE_PRICE_METRIC = 1 only; '
              'AED/sqft. Record-level averages/medians retain outliers and repeated IDs.')
GRAIN_NOTE = ('Transaction numbers are not always unique at row level. Total market value '
              'is intentionally excluded pending transaction-grain reconciliation.')

# DICTIONARY and formula-cache helpers are defined below before main() is called.


def require(condition, message):
    """Stop clearly rather than publishing a workbook that failed a check."""
    if not condition:
        raise ValueError(message)


def sha256(path):
    digest = hashlib.sha256()
    with path.open('rb') as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b''):
            digest.update(block)
    return digest.hexdigest()


def load_source():
    """Retain a literal-text copy for sampling, including the source text 'NA'."""
    require(SOURCE.is_file(), f'Place the Stage 6 feature CSV at {SOURCE}.')
    before = sha256(SOURCE)
    for encoding in ('utf-8-sig', 'cp1252'):
        try:
            raw = pd.read_csv(SOURCE, dtype=str, keep_default_na=False, encoding=encoding)
            break
        except UnicodeDecodeError:
            if encoding == 'cp1252':
                raise ValueError('The CSV encoding could not be read safely.') from None
    expected = [entry[0] for entry in DICTIONARY]
    require(raw.columns.tolist() == expected, 'Expected the 41 Stage 6 columns in source order.')
    require(len(raw) >= SAMPLE_SIZE, 'The source needs at least 5,000 records for this sample.')

    # Only this in-memory analysis copy receives numeric/date types. No cleaning
    # of names or missing categories is performed, and the source is never saved.
    data = raw.replace('', pd.NA).copy()
    numeric = ['TRANS_VALUE', 'PROCEDURE_AREA', 'ACTUAL_AREA', 'TOTAL_BUYER', 'TOTAL_SELLER',
               'TRANSACTION_YEAR', 'TRANSACTION_MONTH', 'TRANSACTION_DAY', 'TRANSACTION_HOUR',
               'IS_SALE_TRANSACTION', 'IS_RESIDENTIAL_FLAG', 'IS_OFFPLAN_FLAG',
               'IS_FREEHOLD_FLAG', 'PROPERTY_SIZE_SQFT', 'VALID_SALE_PRICE_METRIC',
               'SALE_PRICE_PER_SQM', 'SALE_PRICE_PER_SQFT', 'BEDROOM_COUNT']
    for name in numeric:
        converted = pd.to_numeric(data[name], errors='coerce')
        require(not (data[name].notna() & converted.isna()).any(), f'Invalid numeric input: {name}.')
        require(converted.dropna().map(math.isfinite).all(), f'Non-finite numeric input: {name}.')
        data[name] = converted
    data['INSTANCE_DATE'] = pd.to_datetime(data['INSTANCE_DATE'], errors='coerce')
    require(data['INSTANCE_DATE'].notna().all(), 'Date validation failed; no dates will be guessed.')
    require((data['INSTANCE_DATE'].dt.year == data['TRANSACTION_YEAR']).all()
            and (data['INSTANCE_DATE'].dt.month == data['TRANSACTION_MONTH']).all(),
            'Existing date features do not match INSTANCE_DATE.')
    periods = sorted(set(zip(data['TRANSACTION_YEAR'], data['TRANSACTION_MONTH'])))
    require(periods == [(2026, month) for month in range(1, 10)], 'Expected January–September 2026.')
    require(data['INSTANCE_DATE'].max().date().isoformat() == '2026-09-21',
            'Coverage changed. Review the partial-month labels before rebuilding.')
    sales = data.loc[data['GROUP_EN'].eq('Sales').fillna(False)].copy()
    valid = sales.loc[sales['VALID_SALE_PRICE_METRIC'].eq(1)]
    require(len(sales) > 0, 'No Sales records are available.')
    require((valid['TRANS_VALUE'] > 0).all() and (valid['ACTUAL_AREA'] > 0).all()
            and valid['SALE_PRICE_PER_SQFT'].notna().all(), 'Valid Sales price flags are inconsistent.')
    return raw, data, sales, valid, before, encoding


def deterministic_sample(raw):
    """Proportional month/group sampling using the largest-remainder method.

    Hash-based ordering avoids dependence on a library's random-number version.
    Original row positions and a fixed seed determine the same sample each run.
    """
    groups = list(raw.groupby(['TRANSACTION_YEAR', 'TRANSACTION_MONTH', 'GROUP_EN'], sort=True))
    quotas = [len(group) * SAMPLE_SIZE / len(raw) for _, group in groups]
    allocations = [math.floor(quota) for quota in quotas]
    remainder = SAMPLE_SIZE - sum(allocations)
    priority = sorted(range(len(groups)), key=lambda i: (-(quotas[i] - allocations[i]), groups[i][0]))
    for i in priority[:remainder]:
        allocations[i] += 1
    require(all(allocations), 'A source stratum would be omitted; review the sampling design.')
    selected = []
    for (_, group), size in zip(groups, allocations):
        ranked = sorted(group.index, key=lambda i: hashlib.sha256(f'{SAMPLE_SEED}:{i}'.encode()).digest())
        selected.extend(ranked[:size])
    selected.sort()
    require(len(selected) == SAMPLE_SIZE and len(set(selected)) == SAMPLE_SIZE,
            'The sample does not contain exactly 5,000 distinct source rows.')
    return raw.loc[selected], allocations


def native(value):
    """Translate pandas missing values to blank Excel cells, without inventing data."""
    if pd.isna(value):
        return None
    return value.item() if hasattr(value, 'item') else value


def set_freeze_view(ws, coordinate=None):
    """Replace the complete pane/selection state rather than accumulating it.

    openpyxl's freeze_panes setter retains existing selections. Changing B8 to
    A6 used to leave duplicate and nonexistent pane selections, which Excel
    repaired on opening. A single selection in the active pane is sufficient.
    """
    view = ws.sheet_view
    view.pane = None
    view.selection = [Selection(activeCell='A1', sqref='A1')]
    view.view = 'normal'
    view.topLeftCell = 'A1'
    if coordinate and coordinate != 'A1':
        ws.freeze_panes = coordinate
        view.selection = [Selection(pane=view.pane.activePane,
                                    activeCell=coordinate, sqref=coordinate)]


def title(ws, heading, subtitle, last_col):
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_properties.outlinePr.summaryRight = False
    ws.sheet_properties.tabColor = TEAL
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=last_col)
    ws.cell(1, 1, heading).font = Font(name='Aptos Display', size=22, bold=True, color='FFFFFF')
    ws.cell(1, 1).fill = PatternFill('solid', fgColor=NAVY)
    ws.cell(1, 1).alignment = Alignment(vertical='center', indent=1)
    for row in ws.iter_rows(min_row=1, max_row=2, max_col=last_col):
        for cell in row:
            cell.fill = PatternFill('solid', fgColor=NAVY)
    ws.row_dimensions[1].height = 23
    ws.row_dimensions[2].height = 23
    note(ws, 3, subtitle, last_col, height=32)
    set_freeze_view(ws, 'B8')
    ws.sheet_format.defaultRowHeight = 22
    ws.print_options.horizontalCentered = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.oddFooter.center.text = 'Dubai Real Estate Intelligence Platform | Stage 9'
    ws.oddFooter.right.text = 'Page &P of &N'


def note(ws, row, text, last_col, height=34, fill=None):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    cell = ws.cell(row, 1, text)
    cell.font = Font(name='Aptos', size=11, color=MUTED)
    cell.alignment = Alignment(wrap_text=True, vertical='center', indent=1)
    if fill:
        cell.fill = PatternFill('solid', fgColor=fill)
    ws.row_dimensions[row].height = height


def table(ws, row, headers, rows, name, widths=None, formats=None):
    """Create a native filterable Excel table; formulas remain real formulas."""
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row, col, header)
        cell.font = Font(name='Aptos', size=10, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor=NAVY)
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[row].height = 56
    for offset, values in enumerate(rows, row + 1):
        for col, value in enumerate(values, 1):
            cell = ws.cell(offset, col, native(value))
            cell.font = Font(name='Aptos', size=11, color=INK)
            cell.alignment = Alignment(vertical='center')
            cell.border = Border(bottom=Side(style='hair', color=LINE))
            if formats and col in formats:
                cell.number_format = formats[col]
    end = row + len(rows)
    require(end > row, f'{name} would be an empty table.')
    excel_table = Table(displayName=name, ref=f'A{row}:{get_column_letter(len(headers))}{end}')
    excel_table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(excel_table)
    if widths:
        for col, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(col)].width = width
    return row + 1, end


def groups(frame, column, order=None, limit=None, nonnull=False):
    subset = frame.loc[frame[column].notna()] if nonnull else frame
    grouped = list(subset.groupby(column, dropna=False, sort=False))
    if order:
        grouped.sort(key=lambda item: (order.index(item[0]) if item[0] in order else len(order), str(item[0])))
    else:
        grouped.sort(key=lambda item: (-len(item[1]), str(item[0])))
    return grouped[:limit] if limit else grouped


def price_values(frame):
    return frame.loc[frame['VALID_SALE_PRICE_METRIC'].eq(1), 'SALE_PRICE_PER_SQFT']


def build_notes(wb, data, sales, valid, sample):
    ws = wb['Notes']
    title(ws, 'Notes & methodology', 'Official source, analytical scope and reusable KPI inputs', 2)
    ws.column_dimensions['A'].width = 43
    ws.column_dimensions['B'].width = 115
    ids = data.loc[data['TRANSACTION_NUMBER'].notna()].groupby('TRANSACTION_NUMBER')
    repeated = ids.size().loc[lambda counts: counts > 1].index
    grain = data.loc[data['TRANSACTION_NUMBER'].isin(repeated)].groupby('TRANSACTION_NUMBER')['TRANS_VALUE'].nunique()
    metrics = [
        ('Project', 'Dubai Real Estate Intelligence Platform'),
        ('Data source', 'Dubai Land Department official real estate transaction data'),
        ('Coverage', f"{data['INSTANCE_DATE'].min():%d %B %Y} through {data['INSTANCE_DATE'].max():%d %B %Y}"),
        ('September', 'Partial month'),
        ('Source records after cleaning', len(data)),
        ('Source columns', len(data.columns)),
        ('Sales records', len(sales)),
        ('Distinct Sales transaction numbers', sales['TRANSACTION_NUMBER'].nunique()),
        ('Median valid Sale Price per Sqft', valid['SALE_PRICE_PER_SQFT'].median()),
        ('Valid Sale Price Metric records', len(valid)),
        ('Non-Sales records', len(data) - len(sales)),
        ('Distinct overall transaction numbers', data['TRANSACTION_NUMBER'].nunique()),
        ('Repeated transaction numbers', len(repeated)),
        ('Repeated IDs with one distinct value', int(grain.eq(1).sum())),
        ('Repeated IDs with multiple values', int(grain.gt(1).sum())),
        ('Inspection sample rows', len(sample)),
    ]
    refs = {}
    for row, (label, value) in enumerate(metrics, 6):
        ws.cell(row, 1, label).font = Font(name='Aptos', bold=True, color=INK, size=11)
        cell = ws.cell(row, 2, native(value))
        cell.font = Font(name='Aptos', color=INK, size=11)
        cell.number_format = AED_FMT if 'Median' in label else COUNT_FMT
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        refs[label] = f"'Notes'!B{row}"
    notes = [
        '1. Exact duplicate source records were removed during Stage 5.',
        '2. Repeated TRANSACTION_NUMBER values were retained because they can represent multiple property/unit lines. Distinct counts across groups or months are not additive.',
        f'3. The Stage 8 grain audit is reproduced above: {len(repeated):,} repeated IDs; {int(grain.eq(1).sum()):,} with one distinct value and {int(grain.gt(1).sum()):,} with multiple values.',
        '4. The executive summary intentionally excludes a naive Total Sales Value KPI. Neither summing record values nor arbitrarily selecting one per ID establishes economic market value.',
        '5. SALE_PRICE_PER_SQFT is calculated only for Sales with positive transaction value and property area. This workbook additionally requires VALID_SALE_PRICE_METRIC = 1 for every price statistic.',
        '6. No price-per-area outliers were removed or capped. Averages and medians describe record distributions; repeated IDs retain their record weights.',
        '7. Official DLD spelling and capitalization are preserved. BUSINESS BAY and Business Bay remain separate labels. Semantic normalization is deferred.',
        '8. Full data remains in the processed CSV and MySQL pipeline. All workbook summaries use the full CSV; the inspection sample does not drive any KPI.',
        f'9. Sample: proportional allocation by year/month and GROUP_EN, largest-remainder rounding, seed {SAMPLE_SEED}. Within each stratum, source row positions are ordered by SHA-256. No sample weights are applied to summaries.',
        '10. Sample precision: Excel supports about 15 significant numeric digits. Values exceeding that precision are stored as their original CSV text to prevent silent rounding. Other numeric values and dates use native Excel types. Blank source values remain blank; literal NA remains text.',
        '11. Percentages use all Sales records unless stated otherwise. Bedroom analysis includes only Residential Sales with a known bedroom count; the subtype table excludes missing subtypes.',
        '12. September data ends on 21 September 2026. Its observed MoM growth is a partial-month comparison; no extrapolation or forecast is made.',
        '13. Formulas reference only this workbook. Cached results are evaluated from those formulas for immediate previews; Excel recalculates on opening. Re-run the Python builder to refresh from a changed CSV.',
        '14. AED identifies monetary values; property sizes use square metres or square feet as labelled. Top-area lists are descriptive record counts, not investment recommendations.',
    ]
    for row, text in enumerate(notes, 24):
        note(ws, row, text, 2, height=52 if len(text) > 230 else 40, fill=LIGHT if row % 2 == 0 else None)
    set_freeze_view(ws, 'B6')
    ws.print_area = f'A1:B{23 + len(notes)}'
    return refs


def build_monthly(wb, sales):
    ws = wb['Monthly Analysis']
    title(ws, 'Monthly analysis', 'Sales records and distinct identifiers | January–September 2026', 11)
    note(ws, 4, 'September is partial through 21 September. MoM compares observed records and must not be read as full-month growth.', 11, fill='FFF5DF')
    rows = []
    for month in range(1, 10):
        part = sales.loc[sales['TRANSACTION_MONTH'].eq(month)]
        row = 7 + month
        rows.append([2026, month, calendar.month_name[month] + (' (partial)' if month == 9 else ''),
                     len(part), part['TRANSACTION_NUMBER'].nunique(), part['IS_OFFPLAN_EN'].eq('Off-Plan').sum(),
                     part['IS_OFFPLAN_EN'].eq('Ready').sum(), part['USAGE_EN'].eq('Residential').sum(),
                     f'=IFERROR(F{row}/D{row},"")', None if month == 1 else f'=IFERROR(D{row}/D{row-1}-1,"")',
                     'Yes' if month == 9 else 'No'])
    table(ws, 7, ['TRANSACTION_YEAR', 'TRANSACTION_MONTH', 'MONTH', 'SALE_RECORDS',
          'DISTINCT_SALE_TRANSACTION_NUMBERS', 'OFFPLAN_SALE_RECORDS', 'READY_SALE_RECORDS',
          'RESIDENTIAL_SALE_RECORDS', 'OFFPLAN_SHARE', 'MOM_SALE_RECORD_GROWTH', 'IS_PARTIAL_MONTH'],
          rows, 'MonthlySales', [18, 19, 25, 19, 32, 23, 23, 27, 20, 26, 22],
          {**{i: COUNT_FMT for i in range(1, 9) if i != 3}, 9: PCT_FMT, 10: PCT_FMT})
    ws.conditional_formatting.add('J9:J16', ColorScaleRule(start_type='min', start_color='F1DAD5',
        mid_type='num', mid_value=0, mid_color='FFFFFF', end_type='max', end_color='BCDCD6'))
    for cell in ws[16]:
        cell.fill = PatternFill('solid', fgColor='FFF5DF')
    chart = bar_chart(ws, 8, 16, 3, 4, 'Monthly Sales Records · September partial')
    chart.width, chart.height = 29, 11
    chart.series[0].data_points = [DataPoint(idx=8, spPr=GraphicalProperties(solidFill=GOLD))]
    ws.add_chart(chart, 'A20')
    ws.print_area = 'A1:K43'


def build_areas(wb, sales):
    ws = wb['Area Analysis']
    title(ws, 'Area analysis', 'Top 25 source AREA_EN labels by Sales records | Full dataset', 8)
    note(ws, 4, 'Source labels are preserved exactly: BUSINESS BAY and Business Bay remain separate. Semantic normalization can be handled later.', 8, fill=LIGHT)
    note(ws, 5, PRICE_NOTE, 8)
    rows = []
    for offset, (label, part) in enumerate(groups(sales, 'AREA_EN', limit=25, nonnull=True), 8):
        prices = price_values(part)
        rows.append([label, len(part), part['TRANSACTION_NUMBER'].nunique(), part['IS_OFFPLAN_EN'].eq('Off-Plan').sum(),
                     part['IS_OFFPLAN_EN'].eq('Ready').sum(), f'=IFERROR(D{offset}/B{offset},"")', prices.mean(), prices.median()])
    require(len(rows) == 25, 'Fewer than 25 named areas are available.')
    table(ws, 7, ['AREA_EN', 'SALE_RECORDS', 'DISTINCT_SALE_TRANSACTION_NUMBERS', 'OFFPLAN_SALE_RECORDS',
          'READY_SALE_RECORDS', 'OFFPLAN_SHARE', 'AVERAGE_SALE_PRICE_PER_SQFT', 'MEDIAN_SALE_PRICE_PER_SQFT'],
          rows, 'AreaSales', [37, 18, 31, 23, 23, 19, 30, 30],
          {2: COUNT_FMT, 3: COUNT_FMT, 4: COUNT_FMT, 5: COUNT_FMT, 6: PCT_FMT, 7: AED_FMT, 8: AED_FMT})
    ws.conditional_formatting.add('H8:H32', ColorScaleRule(start_type='min', start_color='EFF5F7',
        end_type='max', end_color='9CCBC7'))
    chart = bar_chart(ws, 8, 17, 1, 2, 'Top 10 Areas by Sales Records', horizontal=True)
    chart.width, chart.height = 30, 14
    ws.add_chart(chart, 'A36')
    ws.print_area = 'A1:H65'


def build_segments(wb, sales):
    ws = wb['Segment Analysis']
    title(ws, 'Segment analysis', 'Sales-only record distributions | Distinct transaction counts shown separately', 6)
    note(ws, 4, PRICE_NOTE, 6, height=42)
    specs = [
        ('A. OFF-PLAN VS READY', 'IS_OFFPLAN_EN', sales, None, None, False, 'offplan'),
        ('B. PROPERTY TYPE', 'PROP_TYPE_EN', sales, None, None, False, 'type'),
        ('C. PROPERTY SUBTYPE · TOP 15 NON-NULL', 'PROP_SB_TYPE_EN', sales, None, 15, True, 'subtype'),
        ('D. TRANSACTION VALUE BAND', 'TRANSACTION_VALUE_BAND', sales, VALUE_BANDS, None, False, 'value'),
        ('E. PROPERTY SIZE BAND', 'PROPERTY_SIZE_BAND', sales, SIZE_BANDS, None, False, 'size'),
        ('F. BEDROOM ANALYSIS · RESIDENTIAL SALES', 'BEDROOM_COUNT', sales.loc[sales['USAGE_EN'].eq('Residential') & sales['BEDROOM_COUNT'].notna()], list(range(8)), None, True, 'bedroom'),
        ('G. FREEHOLD ANALYSIS', 'IS_FREE_HOLD_EN', sales, None, None, False, 'freehold'),
    ]
    ranges = {}
    row = 6
    for label, field, frame, ordering, limit, nonnull, kind in specs:
        note(ws, row, label, 6, height=28, fill=LIGHT)
        header = row + 1
        entries = groups(frame, field, order=ordering, limit=limit, nonnull=nonnull)
        rows = []
        for r, (value, part) in enumerate(entries, header + 1):
            percentage = f'=IFERROR(B{r}/SUM(\'Monthly Analysis\'!D8:D16),"")'
            common = [value, len(part)]
            prices = price_values(part)
            distinct = part['TRANSACTION_NUMBER'].nunique()
            if kind == 'offplan':
                values = common + [percentage, prices.mean(), prices.median(), distinct]
                headers = [field, 'SALE_RECORDS', 'PERCENTAGE_OF_SALES', 'AVERAGE_SALE_PRICE_PER_SQFT', 'MEDIAN_SALE_PRICE_PER_SQFT', 'DISTINCT_SALE_TRANSACTION_NUMBERS']
                formats = {2: COUNT_FMT, 3: PCT_FMT, 4: AED_FMT, 5: AED_FMT, 6: COUNT_FMT}
            elif kind == 'value':
                values = common + [percentage, distinct]
                headers = [field, 'SALE_RECORDS', 'PERCENTAGE', 'DISTINCT_SALE_TRANSACTION_NUMBERS']
                formats = {2: COUNT_FMT, 3: PCT_FMT, 4: COUNT_FMT}
            elif kind == 'bedroom':
                values = common + [part['TRANS_VALUE'].mean(), part['PROPERTY_SIZE_SQFT'].mean(), prices.mean(), distinct]
                headers = [field, 'SALE_RECORDS', 'AVERAGE_TRANS_VALUE', 'AVERAGE_PROPERTY_SIZE_SQFT', 'AVERAGE_SALE_PRICE_PER_SQFT', 'DISTINCT_SALE_TRANSACTION_NUMBERS']
                formats = {1: '0', 2: COUNT_FMT, 3: VALUE_FMT, 4: '#,##0.00', 5: AED_FMT, 6: COUNT_FMT}
            elif kind == 'subtype':
                values = common + [distinct, part['TRANS_VALUE'].mean(), prices.mean()]
                headers = [field, 'SALE_RECORDS', 'DISTINCT_SALE_TRANSACTION_NUMBERS', 'AVERAGE_TRANS_VALUE', 'AVERAGE_SALE_PRICE_PER_SQFT']
                formats = {2: COUNT_FMT, 3: COUNT_FMT, 4: VALUE_FMT, 5: AED_FMT}
            else:
                values = common + [percentage, part['TRANS_VALUE'].mean(), prices.mean(), distinct]
                headers = [field, 'SALE_RECORDS', 'PERCENTAGE_OF_SALES', 'AVERAGE_TRANS_VALUE', 'AVERAGE_SALE_PRICE_PER_SQFT', 'DISTINCT_SALE_TRANSACTION_NUMBERS']
                formats = {2: COUNT_FMT, 3: PCT_FMT, 4: VALUE_FMT, 5: AED_FMT, 6: COUNT_FMT}
            rows.append(values)
        start, end = table(ws, header, headers, rows, f'Segment{kind.title()}', [36, 20, 30, 30, 32, 34], formats)
        ranges[kind] = (start, end)
        row = end + 3
    note(ws, row, 'Average TRANS_VALUE is an average source-record value in AED, not a transaction-grain total. No categories are silently merged or renamed.', 6, height=44)
    ws.print_area = f'A1:F{row}'
    return ranges


def sample_value(text, field):
    """Preserve text and precision while using native Excel dates/numbers safely."""
    if text == '':
        return None
    if field in ('INSTANCE_DATE', 'TRANSACTION_DATE'):
        date = datetime.fromisoformat(text)
        return date.date() if field == 'TRANSACTION_DATE' else date
    numeric = {'TRANS_VALUE', 'PROCEDURE_AREA', 'ACTUAL_AREA', 'TOTAL_BUYER', 'TOTAL_SELLER',
               'TRANSACTION_YEAR', 'TRANSACTION_MONTH', 'TRANSACTION_DAY', 'TRANSACTION_HOUR',
               'IS_SALE_TRANSACTION', 'IS_RESIDENTIAL_FLAG', 'IS_OFFPLAN_FLAG', 'IS_FREEHOLD_FLAG',
               'PROPERTY_SIZE_SQFT', 'VALID_SALE_PRICE_METRIC', 'SALE_PRICE_PER_SQM',
               'SALE_PRICE_PER_SQFT', 'BEDROOM_COUNT'}
    if field in numeric:
        number = Decimal(text)
        # CSV decimals outside Excel's 15-digit precision stay verbatim text.
        if Decimal(format(float(number), '.15g')) == number:
            return int(number) if number == number.to_integral() else float(number)
    return text


def build_sample(wb, raw, sample):
    ws = wb['Data Sample']
    title(ws, 'Data sample', 'Deterministic inspection sample | Original 41 source columns', 41)
    note(ws, 4, f'Full analysis-ready dataset contains {len(raw):,} rows and is maintained in the project CSV/MySQL pipeline. This sheet contains a {len(sample):,}-row sample for workbook inspection.', 12, height=38, fill=LIGHT)
    note(ws, 5, 'All summaries use the full dataset. Long-precision decimals remain original text to avoid Excel rounding; see Notes. Dates retain source timestamps.', 12, height=34)
    rows = [[sample_value(text, field) for field, text in zip(raw.columns, values)]
            for values in sample.itertuples(index=False, name=None)]
    formats = {}
    for col, field in enumerate(raw.columns, 1):
        if field == 'INSTANCE_DATE':
            fmt = DATETIME_FMT
        elif field == 'TRANSACTION_DATE':
            fmt = DATE_FMT
        elif field == 'TRANS_VALUE':
            fmt = VALUE_FMT
        elif field in ('SALE_PRICE_PER_SQM', 'SALE_PRICE_PER_SQFT'):
            fmt = AED_FMT
        elif field in ('ACTUAL_AREA', 'PROCEDURE_AREA', 'PROPERTY_SIZE_SQFT'):
            fmt = '#,##0.00'
        else:
            fmt = COUNT_FMT
        formats[col] = fmt
    table(ws, 7, raw.columns.tolist(), rows, 'SourceSample', [24] * 41, formats)
    for name in ('D', 'H', 'J', 'U', 'V'):
        ws.column_dimensions[name].width = 34
    ws.column_dimensions['B'].width = 25
    # A source string beginning '=' is literal data, never a formula to execute.
    for row in ws.iter_rows(min_row=8, max_row=5007, max_col=41):
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = 's'
    set_freeze_view(ws, 'C8')
    ws.print_title_rows = '1:7'


def build_dictionary(wb):
    ws = wb['Data Dictionary']
    title(ws, 'Data dictionary', '22 original DLD fields + 8 Stage 5 date features + 11 Stage 6 analytical features', 5)
    note(ws, 4, 'All 41 CSV fields are documented. ROW_ID is database-only and is not part of this workbook.', 5, fill=LIGHT)
    table(ws, 7, ['FIELD', 'CATEGORY', 'DESCRIPTION', 'SOURCE_OR_DERIVED', 'NOTES'],
          DICTIONARY, 'FieldDictionary', [33, 27, 62, 24, 89])
    for row in ws.iter_rows(min_row=8, max_row=48, max_col=5):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[row[0].row].height = 76
    set_freeze_view(ws, 'C8')
    ws.print_title_rows = '1:7'
    ws.print_area = 'A1:E48'


def chart_font(size=10):
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(sz=size * 100)),
                                endParaRPr=CharacterProperties(sz=size * 100))])


def bar_chart(ws, first, last, label_col, value_col, heading, horizontal=False):
    chart = BarChart()
    chart.type = 'bar' if horizontal else 'col'
    chart.style = 10
    chart.title = heading
    # openpyxl's x_axis is the CATEGORY axis even for horizontal bars;
    # y_axis is the NUMERIC axis. Keep titles attached to their data roles.
    chart.y_axis.title = 'Sales records'
    chart.x_axis.title = 'Source area label' if horizontal else 'Month'
    chart.legend = None
    chart.height, chart.width = 10, 17.5
    chart.add_data(Reference(ws, min_col=value_col, min_row=first - 1, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=label_col, min_row=first, max_row=last))
    chart.series[0].graphicalProperties.solidFill = TEAL
    chart.series[0].graphicalProperties.line.noFill = True
    chart.x_axis.txPr = chart_font(9)
    chart.y_axis.txPr = chart_font(9)
    chart.y_axis.numFmt = COUNT_FMT
    chart.x_axis.numFmt = 'General'
    if horizontal:
        chart.x_axis.scaling.orientation = 'maxMin'
        chart.y_axis.crosses = 'max'
        chart.x_axis.axPos = 'l'
        chart.y_axis.axPos = 'b'
    chart.graphical_properties = GraphicalProperties(solidFill='FFFFFF')
    return chart


def build_dashboard(wb, refs, ranges):
    ws = wb['Executive Summary']
    title(ws, 'Dubai Real Estate Intelligence Platform', 'Dubai Land Department | 2026 Year-to-Date', 16)
    for col in range(1, 17):
        ws.column_dimensions[get_column_letter(col)].width = 11.5
    note(ws, 4, 'Data through 21 September 2026. September is a partial month.', 16, height=30, fill='FFF5DF')
    cards = [
        ('Sales Records', "=SUM('Monthly Analysis'!D8:D16)", COUNT_FMT, 'Imported Sales rows'),
        ('Distinct Sales Transaction Numbers', '=' + refs['Distinct Sales transaction numbers'], COUNT_FMT, 'Distinct IDs across all Sales'),
        ('Median Sale Price per Sqft', '=' + refs['Median valid Sale Price per Sqft'], AED_FMT, 'Valid Sales only · AED/sqft'),
        ('Off-Plan Share of Sales Records', '=IFERROR(SUM(\'Monthly Analysis\'!F8:F16)/SUM(\'Monthly Analysis\'!D8:D16),"")', PCT_FMT, 'Share of all Sales records'),
        ('Ready Share of Sales Records', '=IFERROR(SUM(\'Monthly Analysis\'!G8:G16)/SUM(\'Monthly Analysis\'!D8:D16),"")', PCT_FMT, 'Share of all Sales records'),
        ('Residential Share of Sales Records', '=IFERROR(SUM(\'Monthly Analysis\'!H8:H16)/SUM(\'Monthly Analysis\'!D8:D16),"")', PCT_FMT, 'Share of all Sales records'),
        ('Top Area by Sales Records', "='Area Analysis'!A8", 'General', 'Official source label · records'),
        ('Valid Sale Price Metric Records', '=' + refs['Valid Sale Price Metric records'], COUNT_FMT, 'Sales with eligible size and value'),
    ]
    cells = {}
    for i, (label, formula, fmt, detail) in enumerate(cards):
        left, top = (i % 4) * 4 + 1, 6 + (i // 4) * 6
        right = left + 3
        for row in range(top, top + 5):
            for col in range(left, right + 1):
                ws.cell(row, col).fill = PatternFill('solid', fgColor=LIGHT)
        for first, last in [(top, top + 1), (top + 2, top + 3), (top + 4, top + 4)]:
            ws.merge_cells(start_row=first, start_column=left, end_row=last, end_column=right)
        label_cell = ws.cell(top, left, label)
        label_cell.font = Font(name='Aptos', size=11, bold=True, color=INK)
        label_cell.alignment = Alignment(wrap_text=True, indent=1, vertical='center')
        value = ws.cell(top + 2, left, formula)
        value.font = Font(name='Aptos Display', size=19 if i == 6 else 26, bold=True, color=TEAL)
        value.alignment = Alignment(wrap_text=True, indent=1, vertical='center')
        value.number_format = fmt
        detail_cell = ws.cell(top + 4, left, detail)
        detail_cell.font = Font(name='Aptos', size=9, color=MUTED)
        detail_cell.alignment = Alignment(indent=1, wrap_text=True)
        cells[label] = value.coordinate
    note(ws, 18, GRAIN_NOTE, 16, height=36, fill='FFF5DF')
    note(ws, 19, 'All summaries use the full dataset. Records and distinct transaction numbers are different measures.', 16, height=27)
    monthly = bar_chart(wb['Monthly Analysis'], 8, 16, 3, 4, 'Monthly Sales Records · September partial')
    monthly.series[0].data_points = [DataPoint(idx=8, spPr=GraphicalProperties(solidFill=GOLD))]
    ws.add_chart(monthly, 'A21')
    segment = wb['Segment Analysis']
    start, end = ranges['offplan']
    pie = DoughnutChart()
    pie.title = 'Off-Plan vs Ready Sales · records'
    pie.holeSize = 68
    pie.firstSliceAng = 270
    pie.add_data(Reference(segment, min_col=2, min_row=start - 1, max_row=end), titles_from_data=True)
    pie.set_categories(Reference(segment, min_col=1, min_row=start, max_row=end))
    pie.series[0].data_points = [DataPoint(idx=i, spPr=GraphicalProperties(solidFill=color))
                               for i, color in enumerate([TEAL, NAVY, GOLD])]
    pie.dataLabels = DataLabelList()
    pie.dataLabels.showPercent = True
    pie.dataLabels.showVal = False
    pie.dataLabels.showLeaderLines = True
    pie.legend.position = 'b'
    pie.width, pie.height = 17.5, 10
    ws.add_chart(pie, 'I21')
    area_chart = bar_chart(wb['Area Analysis'], 8, 17, 1, 2, 'Top 10 Areas by Sales Records', horizontal=True)
    area_chart.height = 12
    ws.add_chart(area_chart, 'A36')
    start, end = ranges['type']
    type_chart = bar_chart(segment, start, end, 1, 2, 'Property Type Sales Mix · records')
    type_chart.x_axis.title = 'Property type'
    type_chart.height = 12
    ws.add_chart(type_chart, 'I36')
    note(ws, 54, 'Source: Dubai Land Department | AED monetary values | Source labels preserved | No price outliers removed', 16, height=28)
    # The dashboard does not need frozen panes. Select only its top-left cell.
    set_freeze_view(ws)
    ws.sheet_view.zoomScale = 75
    ws.page_setup.fitToHeight = 1
    ws.print_area = 'A1:P54'
    return cells


def cache_charts(wb):
    """Embed chart data caches so charts also render in workbook preview tools."""
    for ws in wb:
        for chart in ws._charts:
            for series in chart.series:
                for reference in (series.val.numRef,):
                    sheet, region = reference.f.split('!')
                    source = wb[sheet.strip("'").replace("''", "'")]
                    cells = source[region.replace('$', '')]
                    values = [cell.value for row in cells for cell in row]
                    reference.numCache = NumData(ptCount=len(values), formatCode=COUNT_FMT,
                        pt=[NumVal(idx=i, v=value) for i, value in enumerate(values) if value is not None])
                cat_ref = series.cat.numRef.f if series.cat.numRef else series.cat.strRef.f
                sheet, region = cat_ref.split('!')
                source = wb[sheet.strip("'").replace("''", "'")]
                values = [cell.value for row in source[region.replace('$', '')] for cell in row]
                series.cat = AxDataSource(strRef=StrRef(f=cat_ref, strCache=StrData(ptCount=len(values),
                    pt=[StrVal(idx=i, v='' if value is None else str(value)) for i, value in enumerate(values)])))


def validate_package_views(path):
    """Check the saved XLSX view XML, including what openpyxl does not reject.

    This reads the package without editing it. Resolve sheet relationships so
    the check confirms which sheet owns sheet1.xml instead of assuming its name.
    """
    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
          'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
    checks = []
    with zipfile.ZipFile(path) as package:
        require(package.testzip() is None, 'The workbook ZIP package is damaged.')
        root = ET.fromstring(package.read('xl/workbook.xml'))
        relations = ET.fromstring(package.read('xl/_rels/workbook.xml.rels'))
        targets = {item.get('Id'): item.get('Target') for item in relations}
        workbook_views = root.findall('s:bookViews/s:workbookView', ns)
        require(len(workbook_views) == 1 and workbook_views[0].get('activeTab') == '0',
                'Expected one workbook window with Executive Summary active.')
        selected_sheets = []
        for sheet in root.find('s:sheets', ns):
            name = sheet.get('name')
            target = targets[sheet.get('{' + ns['r'] + '}id')]
            member = target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/' + target)
            worksheet = ET.fromstring(package.read(member))
            views = worksheet.findall('s:sheetViews/s:sheetView', ns)
            require(len(views) == 1, f'{name}: expected one worksheet view.')
            view = views[0]
            require(view.get('workbookViewId') == '0' and view.get('view') == 'normal',
                    f'{name}: incompatible workbook view reference or mode.')
            require(view.get('topLeftCell') == 'A1' and 10 <= int(view.get('zoomScale', '100')) <= 400,
                    f'{name}: invalid scrolling origin or zoom.')
            if view.get('tabSelected') in ('1', 'true'):
                selected_sheets.append(name)
            panes = view.findall('s:pane', ns)
            selections = view.findall('s:selection', ns)
            require(len(panes) <= 1 and len(selections) == 1,
                    f'{name}: duplicate pane/selection definitions.')
            selection = selections[0]
            cell = selection.get('activeCell')
            selection_range = selection.get('sqref')
            require(cell is not None and selection_range == cell,
                    f'{name}: active cell must match its single-cell selection.')
            column, row, _, _ = range_boundaries(cell)
            if not panes:
                require(selection.get('pane') is None and cell == 'A1',
                        f'{name}: stale selection refers to a nonexistent pane.')
            else:
                pane = panes[0]
                x = float(pane.get('xSplit', '0'))
                y = float(pane.get('ySplit', '0'))
                require(x >= 0 and y >= 0 and x.is_integer() and y.is_integer() and x + y > 0,
                        f'{name}: invalid frozen split.')
                active = 'bottomRight' if x and y else 'topRight' if x else 'bottomLeft'
                require(pane.get('state') == 'frozen' and pane.get('activePane') == active
                        and selection.get('pane') == active, f'{name}: conflicting active pane.')
                require(pane.get('topLeftCell') == cell and column == x + 1 and row == y + 1,
                        f'{name}: selected cell lies outside the active scrollable pane.')
            if name == 'Executive Summary':
                require(member == 'xl/worksheets/sheet1.xml', 'Executive Summary sheet mapping changed.')
                require(not panes, 'Executive Summary must have no frozen/split pane.')
                checks.append('PASS: sheet1.xml maps to Executive Summary; no pane and one A1 selection')
            checks.append(f'PASS: Saved worksheet-view XML is consistent: {name}')
        require(selected_sheets == ['Executive Summary'], 'Only Executive Summary should be selected.')
    checks.append('PASS: Active workbook tab and selected worksheet agree; no conflicting pane selections')
    return checks


def validate_workbook(path, raw, sample, sales, valid, kpi_cells, formula_results):
    """Reopen the actual saved file in both formula and cached-value modes."""
    checks = validate_package_views(path)
    def check(condition, message):
        require(condition, message)
        checks.append('PASS: ' + message)
    book = load_workbook(path, data_only=False)
    values = load_workbook(path, data_only=True)
    check(book.sheetnames == SHEETS, 'Exactly the seven requested sheets exist and reopen successfully')
    check(len(kpi_cells) == 8 and all(book['Executive Summary'][cell].data_type == 'f' for cell in kpi_cells.values()), 'All eight KPI cards reference analysis-sheet formulas')
    check(len(book['Executive Summary']._charts) == 4, 'Executive Summary contains four charts')
    check(sum(len(ws._charts) for ws in book) == 6, 'Six workbook charts including Monthly and Area analysis charts')
    monthly = values['Monthly Analysis']
    check([monthly.cell(row, 2).value for row in range(8, 17)] == list(range(1, 10)), 'Monthly analysis contains January through September')
    check(monthly['K16'].value == 'Yes' and 'partial' in monthly['C16'].value.lower(), 'September is visibly marked partial')
    check(monthly['J8'].value is None, 'January MoM is blank')
    check(book['Area Analysis'].tables['AreaSales'].ref == 'A7:H32', 'Area analysis contains 25 source labels')
    check(book['Data Sample'].tables['SourceSample'].ref == 'A7:AO5007', 'Sample contains exactly 5,000 data rows and 41 columns')
    check([book['Data Sample'].cell(7, c).value for c in range(1, 42)] == raw.columns.tolist(), 'All 41 sample source headers retain original order and names')
    check([book['Data Dictionary'].cell(r, 1).value for r in range(8, 49)] == raw.columns.tolist(), 'Data dictionary documents all 41 source fields without ROW_ID')
    check(not book._external_links, 'Workbook has no external workbook links')
    errors = {'#REF!', '#DIV/0!', '#VALUE!', '#NAME?', '#N/A', '#NUM!', '#SPILL!'}
    for ws in values:
        check(not any(cell.data_type == 'e' or (isinstance(cell.value, str) and cell.value in errors)
                      for row in ws for cell in row), f'No Excel error cells in {ws.title}')
    for (sheet, coordinate), expected in formula_results.items():
        actual = values[sheet][coordinate].value
        if expected == '':
            require(actual in ('', None), f'Cached blank formula mismatch at {sheet}!{coordinate}.')
        elif isinstance(expected, (float, int)):
            require(actual is not None and math.isclose(actual, expected, rel_tol=1e-12, abs_tol=1e-10),
                    f'Cached formula mismatch at {sheet}!{coordinate}.')
        else:
            require(actual == expected, f'Cached text formula mismatch at {sheet}!{coordinate}.')
    check(True, f'All {len(formula_results)} formulas evaluated and cached without broken references')
    summary = values['Executive Summary']
    expected_kpis = [len(sales), sales['TRANSACTION_NUMBER'].nunique(), valid['SALE_PRICE_PER_SQFT'].median(),
                     sales['IS_OFFPLAN_EN'].eq('Off-Plan').sum() / len(sales),
                     sales['IS_OFFPLAN_EN'].eq('Ready').sum() / len(sales),
                     sales['USAGE_EN'].eq('Residential').sum() / len(sales),
                     groups(sales, 'AREA_EN', limit=1, nonnull=True)[0][0], len(valid)]
    for (label, cell), expected in zip(kpi_cells.items(), expected_kpis):
        actual = summary[cell].value
        check(actual == expected if isinstance(expected, str) else math.isclose(actual, expected, rel_tol=1e-12),
              f'KPI matches independent full-source calculation: {label}')
    check(not any(isinstance(cell.value, str) and 'total sales value' in cell.value.lower()
                  for row in book['Executive Summary'] for cell in row), 'No Total Sales Value KPI exists')
    # Compare all 205,000 sample cells with their selected source values, not
    # merely the row count. Exact long decimals remain strings in this check.
    for r, original in enumerate(sample.itertuples(index=False, name=None), 8):
        for c, (field, text) in enumerate(zip(raw.columns, original), 1):
            actual = values['Data Sample'].cell(r, c).value
            expected = sample_value(text, field)
            if field == 'TRANSACTION_DATE' and isinstance(actual, datetime):
                actual = actual.date()
            require(actual == expected, f'Sample value mismatch at source field {field}, workbook row {r}.')
    check(True, 'All 205,000 sample cells preserve their selected source values')
    check(path.stat().st_size < 50 * 1024 * 1024, 'Workbook is below 50 MB')
    book.close()
    values.close()
    return checks, dict(zip(kpi_cells, expected_kpis))


def main():
    try:
        print('Stage 9: reading the protected Stage 6 feature CSV...')
        raw, data, sales, valid, before, encoding = load_source()
        sample, allocations = deterministic_sample(raw)
        book = Workbook()
        book.remove(book.active)
        for name in SHEETS:
            book.create_sheet(name)
        book.active = 0
        book.views[0].activeTab = 0
        book.views[0].firstSheet = 0
        for index, worksheet in enumerate(book):
            worksheet.sheet_view.tabSelected = index == 0
        book.properties.creator = 'Dubai Real Estate Intelligence Platform'
        book.properties.title = 'Dubai Real Estate Intelligence Platform | 2026 YTD'
        book.properties.subject = 'Stage 9 - Excel Executive Analysis'
        book.properties.description = 'Official DLD data; record-grain summaries; September partial.'
        book.calculation = CalcProperties(calcMode='auto', fullCalcOnLoad=True, forceFullCalc=True)
        refs = build_notes(book, data, sales, valid, sample)
        build_monthly(book, sales)
        build_areas(book, sales)
        ranges = build_segments(book, sales)
        build_sample(book, raw, sample)
        build_dictionary(book)
        kpi_cells = build_dashboard(book, refs, ranges)
        cache_charts(book)
        print('Evaluating formulas, preserving cached results and validating the workbook...')
        formula_results = evaluate_formulas(book)
        OUTPUT.parent.mkdir(parents=True, exist_ok=True)
        REPORT.parent.mkdir(parents=True, exist_ok=True)
        # Publish only a successfully validated file; an existing workbook is
        # not replaced by a half-written file if an error occurs.
        with tempfile.TemporaryDirectory(prefix='dld-stage9-') as temporary:
            candidate = Path(temporary) / OUTPUT.name
            book.save(candidate)
            cache_formula_results(candidate, formula_results)
            checks, kpis = validate_workbook(candidate, raw, sample, sales, valid, kpi_cells, formula_results)
            after = sha256(SOURCE)
            require(before == after, 'The source hash changed during processing; workbook not published.')
            checks.append('PASS: Source CSV SHA-256 before and after are identical')
            file_size = candidate.stat().st_size
            report = [
                'Dubai Real Estate Intelligence Platform', 'Stage 9 - Excel Executive Analysis Workbook',
                '=' * 76, 'Status: SUCCESS', f'Input CSV: {SOURCE.relative_to(ROOT)}',
                f'Output workbook: {OUTPUT.relative_to(ROOT)}', f'Source encoding: {encoding}',
                f'Source rows: {len(raw):,}', f'Source columns: {len(raw.columns)}',
                f'Coverage: {data["INSTANCE_DATE"].min()} through {data["INSTANCE_DATE"].max()}',
                'September 2026: PARTIAL MONTH', 'Workbook sheets: ' + ' | '.join(SHEETS),
                f'Sample rows: {len(sample):,}', f'Sample columns: {len(sample.columns)}',
                f'Sample method: proportional month/group strata; largest remainder; SHA-256 seed {SAMPLE_SEED}',
                f'Sample strata: {len(allocations)}; allocation range: {min(allocations)}–{max(allocations)}',
                'Sample source row positions SHA-256: ' + hashlib.sha256(','.join(map(str, sample.index)).encode()).hexdigest(),
                'Chart count: 6 total; 4 Executive Summary; 1 Monthly Analysis; 1 Area Analysis',
                f'Formula cells: {len(formula_results)}; internally evaluated and cached; Excel recalculates on open',
                f'Workbook file size: {file_size:,} bytes ({file_size / 1024 / 1024:.2f} MB)',
                '', 'KPI values calculated from the full source:',
            ]
            for label, value in kpis.items():
                displayed = f'{value:.2%}' if 'Share' in label else (f'AED {value:,.2f}' if label == 'Median Sale Price per Sqft' else f'{value:,}' if isinstance(value, (int, float)) else str(value))
                report.append(f'  {label}: {displayed}')
            report += ['', f'Source SHA-256 before: {before}', f'Source SHA-256 after: {after}',
                       'Source unchanged: YES', '', 'Validation results:', *checks, '',
                       'Excel worksheet-view repair diagnosis and fix:',
                       'Confirmed xl/worksheets/sheet1.xml belongs to Executive Summary.',
                       'Previous cause: repeated freeze_panes assignments retained old selections. B8 then A6 left duplicate bottomLeft selections and a stale bottomRight selection.',
                       'Fix: the generator resets pane and selection metadata before each view configuration. Executive Summary has no freeze pane and one A1 selection; other sheets retain their intended freezes with one compatible selection.',
                       'Saved XML validation covers all seven sheets, active cells, split geometry, selected tabs and workbook view references.',
                       'Native Microsoft Excel repair-free opening: not established by this Python build. A manual Excel reopen is still required unless a separate native validation is recorded.',
                       '',
                       'All summaries use full-source records, with case-sensitive source categories.',
                       'No naive Total Sales Value KPI, outlier removal, recommendations or forecasts.',
                       'Sample decimals exceeding Excel precision remain original text; missing stays blank.',
                       'No MySQL access, credentials, Power BI work or previous-stage modifications.']
            # Copy only completed bytes, then atomically replace within each
            # destination directory (the temporary directory may be another disk).
            output_temp = OUTPUT.with_suffix('.xlsx.tmp')
            output_temp.write_bytes(candidate.read_bytes())
            output_temp.replace(OUTPUT)
            report_temp = REPORT.with_suffix('.txt.tmp')
            report_temp.write_text('\n'.join(report) + '\n', encoding='utf-8')
            report_temp.replace(REPORT)
        print(f'Workbook created: {OUTPUT.relative_to(ROOT)}')
        print(f'Source: {len(raw):,} rows / {len(raw.columns)} columns; sample: {len(sample):,} rows.')
        print(f'Seven sheets, six charts, eight KPI cards. Size: {file_size / 1024 / 1024:.2f} MB.')
        print('Source unchanged: YES. All workbook checks: PASS.')
        print(f'Report: {REPORT.relative_to(ROOT)}')
        return 0
    except (ValueError, OSError, pd.errors.ParserError, zipfile.BadZipFile) as error:
        print(f'Stage 9 stopped: {error}')
        print('Check the source path and close the output workbook before retrying. No database was accessed.')
        return 1


"""Stage 9 data dictionary; descriptions preserve the source field semantics."""

DICTIONARY = [
    ("TRANSACTION_NUMBER", "Original DLD field", "Official transaction identifier supplied by DLD.", "Original DLD", "Not a unique row key. A transaction number may represent several property/unit lines; report record counts separately from distinct transaction numbers."),
    ("INSTANCE_DATE", "Original DLD field", "Transaction date and time recorded in the source.", "Original DLD", "Parsed during Stage 5. Rows were retained if conversion failed; no timezone has been inferred."),
    ("GROUP_EN", "Original DLD field", "English transaction group, such as Sales, Mortgage or Gifts.", "Original DLD", "The exact category Sales defines Sales records in this workbook; official spelling and capitalization are preserved."),
    ("PROCEDURE_EN", "Original DLD field", "English procedure description recorded by DLD.", "Original DLD", "Preserved as supplied; this field does not determine the Sales flag."),
    ("IS_OFFPLAN_EN", "Original DLD field", "Official Off-Plan or Ready category.", "Original DLD", "Exact source labels are retained. Missing or unexpected categories are not guessed."),
    ("IS_FREE_HOLD_EN", "Original DLD field", "Official Free Hold or Non Free Hold category.", "Original DLD", "Preserves the DLD category; missing or unexpected categories are not guessed."),
    ("USAGE_EN", "Original DLD field", "Official English property-use category.", "Original DLD", "Residential analysis uses the exact value Residential; source categories are not normalized."),
    ("AREA_EN", "Original DLD field", "Official English area label attached to the record.", "Original DLD", "Capitalization variants such as BUSINESS BAY and Business Bay remain separate labels; semantic normalization is deferred."),
    ("PROP_TYPE_EN", "Original DLD field", "Official English property-type category.", "Original DLD", "Used for property-type segmentation without changing source categories."),
    ("PROP_SB_TYPE_EN", "Original DLD field", "Official English property-subtype category.", "Original DLD", "Missing subtypes are excluded from the top-15 subtype table and are not filled."),
    ("TRANS_VALUE", "Original DLD field", "Transaction value reported in AED for the source record.", "Original DLD", "Repeated IDs can carry the same or different values. Row-level averages describe records; summing records is not established as total market sales value."),
    ("PROCEDURE_AREA", "Original DLD field", "Procedure area field supplied by DLD.", "Original DLD", "Retained separately from ACTUAL_AREA; not assumed to be the property's size and not used as the price-per-area denominator."),
    ("ACTUAL_AREA", "Original DLD field", "Official DLD property size in square metres.", "Original DLD", "Positive numeric values support property-size and sale-price-per-area features. Missing, zero and negative values are not imputed."),
    ("ROOMS_EN", "Original DLD field", "Official English room or property-category label.", "Original DLD", "Not every category is a bedroom count. Literal source text such as NA is preserved; missing values are not filled."),
    ("PARKING", "Original DLD field", "Parking information recorded in the official source.", "Original DLD", "Preserved as supplied; may contain detailed text and is not assumed to be a numeric parking-space count."),
    ("NEAREST_METRO_EN", "Original DLD field", "English nearest-metro label supplied by DLD.", "Original DLD", "A source location label, not a measured distance or calculated accessibility score."),
    ("NEAREST_MALL_EN", "Original DLD field", "English nearest-mall label supplied by DLD.", "Original DLD", "A source location label, not a measured distance or calculated accessibility score."),
    ("NEAREST_LANDMARK_EN", "Original DLD field", "English nearest-landmark label supplied by DLD.", "Original DLD", "A source location label, not a measured distance or calculated accessibility score."),
    ("TOTAL_BUYER", "Original DLD field", "Buyer-total field supplied by DLD.", "Original DLD", "Earlier profiling found only zeros. Retained unchanged and not treated as a reliable count of actual people."),
    ("TOTAL_SELLER", "Original DLD field", "Seller-total field supplied by DLD.", "Original DLD", "Earlier profiling found only zeros. Retained unchanged and not treated as a reliable count of actual people."),
    ("MASTER_PROJECT_EN", "Original DLD field", "Official English master-project label.", "Original DLD", "Missing values are preserved; no master-project names are inferred."),
    ("PROJECT_EN", "Original DLD field", "Official English project label.", "Original DLD", "Missing names are preserved. Project analysis excludes missing labels without filling or renaming them."),
    ("TRANSACTION_DATE", "Stage 5 date feature", "Calendar date extracted from INSTANCE_DATE.", "Derived - Stage 5", "Contains the date without the time of day; missing when INSTANCE_DATE is missing or unparseable."),
    ("TRANSACTION_YEAR", "Stage 5 date feature", "Calendar year extracted from INSTANCE_DATE.", "Derived - Stage 5", "Missing when INSTANCE_DATE is missing or unparseable."),
    ("TRANSACTION_MONTH", "Stage 5 date feature", "Calendar month number extracted from INSTANCE_DATE, from 1 to 12.", "Derived - Stage 5", "Use with TRANSACTION_YEAR for chronological monthly ordering."),
    ("TRANSACTION_MONTH_NAME", "Stage 5 date feature", "Calendar month name extracted from INSTANCE_DATE.", "Derived - Stage 5", "A display label; sort by year and month number rather than alphabetically."),
    ("TRANSACTION_QUARTER", "Stage 5 date feature", "Calendar-quarter label extracted from INSTANCE_DATE.", "Derived - Stage 5", "Q1, Q2, Q3 or Q4; missing for an unavailable date."),
    ("TRANSACTION_DAY", "Stage 5 date feature", "Day of the calendar month extracted from INSTANCE_DATE.", "Derived - Stage 5", "Ranges from 1 to 31 when populated; this is not the weekday number."),
    ("TRANSACTION_DAY_NAME", "Stage 5 date feature", "Weekday name extracted from INSTANCE_DATE.", "Derived - Stage 5", "Missing when INSTANCE_DATE is missing or unparseable."),
    ("TRANSACTION_HOUR", "Stage 5 date feature", "Hour of day extracted from INSTANCE_DATE.", "Derived - Stage 5", "Uses the source timestamp's hour, from 0 to 23; no timezone conversion has been inferred."),
    ("IS_SALE_TRANSACTION", "Stage 6 analytical feature", "Sales indicator based only on GROUP_EN.", "Derived - Stage 6", "1 when GROUP_EN = Sales; 0 otherwise. Repeated transaction numbers remain separate source records."),
    ("IS_RESIDENTIAL_FLAG", "Stage 6 analytical feature", "Residential-use indicator based on USAGE_EN.", "Derived - Stage 6", "1 when USAGE_EN = Residential; 0 otherwise. Preserves the original usage field."),
    ("IS_OFFPLAN_FLAG", "Stage 6 analytical feature", "Nullable numeric representation of the source Off-Plan/Ready category.", "Derived - Stage 6", "Off-Plan = 1; Ready = 0; missing or unrecognized source categories remain missing."),
    ("IS_FREEHOLD_FLAG", "Stage 6 analytical feature", "Nullable numeric representation of the source freehold category.", "Derived - Stage 6", "Free Hold = 1; Non Free Hold = 0; missing or unrecognized source categories remain missing."),
    ("PROPERTY_SIZE_SQFT", "Stage 6 analytical feature", "Property size in square feet, derived as ACTUAL_AREA multiplied by 10.7639104167.", "Derived - Stage 6", "Calculated only when ACTUAL_AREA is numeric and positive. Applies to eligible records of any transaction group; ACTUAL_AREA is retained."),
    ("VALID_SALE_PRICE_METRIC", "Stage 6 analytical feature", "Eligibility indicator for sale price-per-area metrics.", "Derived - Stage 6", "1 only when GROUP_EN = Sales, TRANS_VALUE > 0 and ACTUAL_AREA > 0; 0 otherwise."),
    ("SALE_PRICE_PER_SQM", "Stage 6 analytical feature", "Sales-only price per square metre in AED, calculated as TRANS_VALUE / ACTUAL_AREA.", "Derived - Stage 6", "Populated only where VALID_SALE_PRICE_METRIC = 1. Non-Sales and invalid records remain missing; no outliers were removed or capped."),
    ("SALE_PRICE_PER_SQFT", "Stage 6 analytical feature", "Sales-only price per square foot in AED, calculated as TRANS_VALUE / PROPERTY_SIZE_SQFT.", "Derived - Stage 6", "Populated only where VALID_SALE_PRICE_METRIC = 1. Workbook price summaries additionally require GROUP_EN = Sales; no outliers were removed or capped."),
    ("TRANSACTION_VALUE_BAND", "Stage 6 analytical feature", "Ordered AED value band based on TRANS_VALUE.", "Derived - Stage 6", "Under 500K; 500K - 1M; 1M - 2M; 2M - 5M; 5M - 10M; 10M+. Lower boundaries are inclusive and upper boundaries exclusive; missing input remains missing."),
    ("PROPERTY_SIZE_BAND", "Stage 6 analytical feature", "Ordered property-size band based on ACTUAL_AREA in square metres.", "Derived - Stage 6", "Under 50 sqm; 50 - 100 sqm; 100 - 200 sqm; 200 - 500 sqm; 500+ sqm. Positive values only; lower boundaries inclusive and upper boundaries exclusive."),
    ("BEDROOM_COUNT", "Stage 6 analytical feature", "Bedroom count mapped from recognized ROOMS_EN categories.", "Derived - Stage 6", "Studio = 0; 1 B/R through 7 B/R = 1 through 7. Office, Shop, PENTHOUSE, missing and all other categories remain missing; no bedroom values are guessed."),
]


"""Evaluate this workbook's small formula language and save Excel value caches.

This is deliberately not a general Excel calculation engine. Unsupported syntax
stops validation instead of guessing a result. It never uses Python eval().
"""

from datetime import date, datetime
from decimal import Decimal
import math
import os
from pathlib import Path
import posixpath
import re
import tempfile
import xml.etree.ElementTree as ET
from zipfile import ZipFile

from openpyxl.formula import Tokenizer
from openpyxl.utils.cell import column_index_from_string, get_column_letter
from openpyxl.utils.datetime import CALENDAR_MAC_1904, CALENDAR_WINDOWS_1900, to_excel


class FormulaValidationError(ValueError):
    """A formula, reference or saved workbook cannot be verified safely."""


class _CellCalculationError(Exception):
    """A calculation error, such as division by zero, which IFERROR may handle."""


class _FormulaParser:
    """Turn openpyxl tokens into a small arithmetic abstract syntax tree."""

    def __init__(self, formula):
        self.tokens = [token for token in Tokenizer(formula).items
                       if token.type != "WHITE-SPACE"]
        self.position = 0

    def peek(self):
        return self.tokens[self.position] if self.position < len(self.tokens) else None

    def take(self):
        token = self.peek()
        if token is None:
            raise FormulaValidationError("Formula ends unexpectedly.")
        self.position += 1
        return token

    def parse(self):
        node = self.expression(0)
        if self.peek() is not None:
            raise FormulaValidationError("Unsupported trailing formula syntax.")
        return node

    def expression(self, minimum_precedence):
        node = self.atom()
        # Multiplication/division bind more tightly than addition/subtraction.
        precedence = {"+": 1, "-": 1, "*": 2, "/": 2}
        while self.peek() is not None:
            token = self.peek()
            if token.type != "OPERATOR-INFIX" or token.value not in precedence:
                break
            level = precedence[token.value]
            if level < minimum_precedence:
                break
            operator = self.take().value
            node = ("binary", operator, node, self.expression(level + 1))
        return node

    def atom(self):
        token = self.take()
        if token.type == "OPERATOR-PREFIX" and token.value in ("+", "-"):
            return ("unary", token.value, self.atom())
        if token.type == "OPERAND":
            if token.subtype == "NUMBER":
                value = float(token.value)
                if not math.isfinite(value):
                    raise FormulaValidationError("Non-finite formula number.")
                return ("literal", value)
            if token.subtype == "TEXT":
                return ("literal", token.value[1:-1].replace('""', '"'))
            if token.subtype == "RANGE":
                return ("reference", token.value)
            raise FormulaValidationError("Unsupported formula operand.")
        if token.type == "PAREN" and token.subtype == "OPEN":
            node = self.expression(0)
            closing = self.take()
            if closing.type != "PAREN" or closing.subtype != "CLOSE":
                raise FormulaValidationError("Unbalanced formula parentheses.")
            return node
        if token.type == "FUNC" and token.subtype == "OPEN":
            name = token.value[:-1].upper()
            if name not in ("SUM", "IFERROR"):
                raise FormulaValidationError(f"Unsupported formula function: {name}.")
            arguments = [self.expression(0)]
            while self.peek() is not None and self.peek().type == "SEP":
                separator = self.take()
                if separator.subtype != "ARG" or separator.value != ",":
                    raise FormulaValidationError("Unsupported argument separator.")
                arguments.append(self.expression(0))
            closing = self.take()
            if closing.type != "FUNC" or closing.subtype != "CLOSE":
                raise FormulaValidationError("Unbalanced function parentheses.")
            if name == "IFERROR" and len(arguments) != 2:
                raise FormulaValidationError("IFERROR requires exactly two arguments.")
            return ("function", name, arguments)
        raise FormulaValidationError("Unsupported formula syntax.")


def _reference_cells(workbook, current_sheet, reference):
    """Resolve local A1 references; reject external books and missing targets."""
    if "[" in reference or "]" in reference:
        raise FormulaValidationError("External workbook references are not allowed.")
    if "!" in reference:
        sheet_name, address = reference.rsplit("!", 1)
        if sheet_name.startswith("'") and sheet_name.endswith("'"):
            sheet_name = sheet_name[1:-1].replace("''", "'")
    else:
        sheet_name, address = current_sheet, reference
    if sheet_name not in workbook.sheetnames:
        raise FormulaValidationError(f"Missing worksheet reference: {sheet_name}.")
    match = re.fullmatch(r"\$?([A-Za-z]{1,3})\$?([1-9][0-9]*)(?::\$?([A-Za-z]{1,3})\$?([1-9][0-9]*))?", address)
    if match is None:
        raise FormulaValidationError(f"Unsupported or invalid cell reference: {reference}.")
    first_column, first_row, last_column, last_row = match.groups()
    start_column = column_index_from_string(first_column.upper())
    end_column = column_index_from_string((last_column or first_column).upper())
    start_row, end_row = int(first_row), int(last_row or first_row)
    if not (1 <= start_column <= end_column <= 16384 and 1 <= start_row <= end_row <= 1048576):
        raise FormulaValidationError(f"Cell reference is outside Excel bounds: {reference}.")
    worksheet = workbook[sheet_name]
    coordinates = []
    for row in range(start_row, end_row + 1):
        for column in range(start_column, end_column + 1):
            coordinate = f"{get_column_letter(column)}{row}"
            # Existing explicitly blank cells are allowed. A never-created cell
            # is treated as a likely typo in this controlled workbook builder.
            if (row, column) not in worksheet._cells:
                raise FormulaValidationError(f"Missing referenced cell: {sheet_name}!{coordinate}.")
            coordinates.append((sheet_name, coordinate))
    return coordinates, ":" in address


def _tree_references(node):
    if node[0] == "reference":
        yield node[1]
    elif node[0] == "unary":
        yield from _tree_references(node[2])
    elif node[0] == "binary":
        yield from _tree_references(node[2])
        yield from _tree_references(node[3])
    elif node[0] == "function":
        for argument in node[2]:
            yield from _tree_references(argument)


def _number(value):
    if value is None or value == "":
        return 0
    if isinstance(value, (int, float, Decimal)) and math.isfinite(float(value)):
        return float(value)
    raise _CellCalculationError("A calculation requires a number.")


def evaluate_formulas(workbook):
    """Calculate every formula from workbook cells and return a cache mapping.

    Keys are (worksheet title, cell coordinate). All references and dependency
    cycles are checked before calculating, including IFERROR fallback branches.
    """
    trees = {}
    for worksheet in workbook.worksheets:
        for cell in worksheet._cells.values():
            if cell.data_type == "f":
                key = (worksheet.title, cell.coordinate)
                try:
                    trees[key] = _FormulaParser(cell.value).parse()
                except Exception as error:
                    raise FormulaValidationError(f"Cannot parse formula at {key[0]}!{key[1]}: {error}") from error

    dependencies = {}
    for key, node in trees.items():
        dependencies[key] = set()
        for reference in _tree_references(node):
            cells, _ = _reference_cells(workbook, key[0], reference)
            dependencies[key].update(target for target in cells if target in trees)

    visiting, complete = set(), set()

    def check_cycles(key):
        if key in visiting:
            raise FormulaValidationError(f"Formula cycle found at {key[0]}!{key[1]}.")
        if key in complete:
            return
        visiting.add(key)
        for dependency in dependencies[key]:
            check_cycles(dependency)
        visiting.remove(key)
        complete.add(key)

    for key in trees:
        check_cycles(key)

    results = {}

    def cell_value(key):
        if key in results:
            return results[key]
        if key not in trees:
            cell = workbook[key[0]][key[1]]
            if cell.data_type == "e":
                raise _CellCalculationError("Referenced cell contains an Excel error.")
            return cell.value
        value = calculate(trees[key], key[0])
        if isinstance(value, list):
            raise FormulaValidationError("A range cannot be the final value of a single-cell formula.")
        if isinstance(value, (float, Decimal)) and not math.isfinite(float(value)):
            raise FormulaValidationError("A formula produced a non-finite number.")
        # Excel treats a direct reference to a blank cell as zero.
        results[key] = 0 if value is None else value
        return results[key]

    def calculate(node, sheet_name):
        kind = node[0]
        if kind == "literal":
            return node[1]
        if kind == "reference":
            cells, is_range = _reference_cells(workbook, sheet_name, node[1])
            values = [cell_value(key) for key in cells]
            return values if is_range else values[0]
        if kind == "unary":
            value = _number(calculate(node[2], sheet_name))
            return value if node[1] == "+" else -value
        if kind == "binary":
            left = _number(calculate(node[2], sheet_name))
            right = _number(calculate(node[3], sheet_name))
            if node[1] == "+":
                return left + right
            if node[1] == "-":
                return left - right
            if node[1] == "*":
                return left * right
            if right == 0:
                raise _CellCalculationError("Division by zero.")
            return left / right
        if node[1] == "IFERROR":
            try:
                return calculate(node[2][0], sheet_name)
            except _CellCalculationError:
                return calculate(node[2][1], sheet_name)
        total = 0.0
        for argument in node[2]:
            value = calculate(argument, sheet_name)
            if isinstance(value, list):
                # Like Excel SUM over a range, skip text and blank source cells.
                for item in value:
                    if isinstance(item, (int, float, Decimal)) and not isinstance(item, bool):
                        total += _number(item)
            elif value is None or value == "":
                continue
            else:
                total += _number(value)
        return total

    for key in trees:
        try:
            cell_value(key)
        except _CellCalculationError as error:
            raise FormulaValidationError(f"Uncaught calculation error at {key[0]}!{key[1]}: {error}") from error
    return results


def cache_formula_results(path, results):
    """Save calculated values inside XLSX XML while retaining every formula.

    openpyxl writes formulas but does not calculate them. These cached values
    make the workbook readable in previews before Excel recalculates it.
    """
    path = Path(path)
    main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    document_relationships = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    package_relationships = "http://schemas.openxmlformats.org/package/2006/relationships"
    ET.register_namespace("", main)
    replacements, seen = {}, set()
    with ZipFile(path, "r") as source:
        workbook_root = ET.fromstring(source.read("xl/workbook.xml"))
        relationships = ET.fromstring(source.read("xl/_rels/workbook.xml.rels"))
        targets = {element.attrib["Id"]: element.attrib["Target"]
                   for element in relationships.findall(f"{{{package_relationships}}}Relationship")}
        properties = workbook_root.find(f"{{{main}}}workbookPr")
        uses_1904 = properties is not None and properties.attrib.get("date1904") in ("1", "true")
        epoch = CALENDAR_MAC_1904 if uses_1904 else CALENDAR_WINDOWS_1900
        for sheet in workbook_root.find(f"{{{main}}}sheets"):
            target = targets[sheet.attrib[f"{{{document_relationships}}}id"]]
            member = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join("xl", target))
            worksheet_root = ET.fromstring(source.read(member))
            changed = False
            for cell in worksheet_root.iter(f"{{{main}}}c"):
                if cell.find(f"{{{main}}}f") is None:
                    continue
                key = (sheet.attrib["name"], cell.attrib["r"])
                if key not in results:
                    raise FormulaValidationError(f"Missing calculated cache for {key[0]}!{key[1]}.")
                value = results[key]
                for child in list(cell):
                    if child.tag in (f"{{{main}}}v", f"{{{main}}}is"):
                        cell.remove(child)
                cached = ET.SubElement(cell, f"{{{main}}}v")
                if isinstance(value, bool):
                    cell.set("t", "b")
                    cached.text = "1" if value else "0"
                elif isinstance(value, str):
                    cell.set("t", "str")
                    cached.text = value
                elif value is None:
                    cell.set("t", "str")
                    cached.text = ""
                elif isinstance(value, (date, datetime)):
                    cell.attrib.pop("t", None)
                    cached.text = repr(to_excel(value, epoch))
                elif isinstance(value, (int, float, Decimal)) and math.isfinite(float(value)):
                    cell.attrib.pop("t", None)
                    cached.text = str(value)
                else:
                    raise FormulaValidationError(f"Unsupported cached value at {key[0]}!{key[1]}.")
                seen.add(key)
                changed = True
            if changed:
                replacements[member] = ET.tostring(worksheet_root, encoding="utf-8", xml_declaration=True)
        if seen != set(results):
            raise FormulaValidationError("Calculated formula keys do not match the saved workbook.")
        # Replace the XLSX atomically, so a failed write cannot leave half a ZIP.
        descriptor, temporary_name = tempfile.mkstemp(prefix="stage9-cache-", suffix=".xlsx", dir=path.parent)
        os.close(descriptor)
        try:
            with ZipFile(temporary_name, "w") as destination:
                for item in source.infolist():
                    destination.writestr(item, replacements.get(item.filename, source.read(item.filename)))
            os.replace(temporary_name, path)
        finally:
            if os.path.exists(temporary_name):
                os.unlink(temporary_name)



if __name__ == '__main__':
    raise SystemExit(main())
